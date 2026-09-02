"""
Peer Comparison Excel Report Exporter.
Generates comprehensive Excel workbooks containing peer group percentile matrices,
rankings, and summary statistics across Nifty 100 sectors.
"""

import os
import logging
from typing import Optional
import pandas as pd

from src.analytics.peer_engine import PeerEngine
from src.analytics.radar_chart import generate_peer_radar_chart

logger = logging.getLogger(__name__)

class PeerExporter:
    """Exports peer group comparative analysis reports to Excel."""

    def __init__(self, peer_engine: Optional[PeerEngine] = None, output_dir: str = "output"):
        self.peer_engine = peer_engine or PeerEngine()
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def export_peer_report(self, file_name: str = "peer_comparison_report.xlsx") -> str:
        """
        Calculates all peer percentiles, generates peer comparison sheets,
        and saves a formatted Excel workbook.
        """
        output_path = os.path.join(self.output_dir, file_name)
        
        # Load and compute percentiles
        df_peers = self.peer_engine.compute_inverse_debt_percentiles()
        self.peer_engine.save_peer_percentiles_to_db(df_peers)

        writer = pd.ExcelWriter(output_path, engine="openpyxl")

        # 1. Master Overall Percentiles Sheet
        cols = [
            "company_id", "company_name", "effective_peer_group", "broad_sector",
            "return_on_equity_pct", "return_on_equity_pct_percentile",
            "roce_pct", "roce_pct_percentile",
            "debt_to_equity", "debt_to_equity_percentile",
            "net_profit_margin_pct", "net_profit_margin_pct_percentile",
            "operating_profit_margin_pct", "operating_profit_margin_pct_percentile",
            "asset_turnover", "asset_turnover_percentile",
            "free_cash_flow_cr", "free_cash_flow_cr_percentile",
            "revenue_cagr_5yr", "revenue_cagr_5yr_percentile",
            "overall_peer_percentile"
        ]
        existing_cols = [c for c in cols if c in df_peers.columns]
        df_peers[existing_cols].to_excel(writer, sheet_name="All Peer Percentiles", index=False)

        # 2. Sector-by-Sector Tabs (Top 5 major peer groups)
        top_groups = df_peers["effective_peer_group"].value_counts().head(5).index
        for grp in top_groups:
            df_grp = df_peers[df_peers["effective_peer_group"] == grp].copy()
            df_grp = df_grp.sort_values(by="overall_peer_percentile", ascending=False)
            sheet_name = str(grp)[:30].replace("/", "_")
            df_grp[existing_cols].to_excel(writer, sheet_name=sheet_name, index=False)

        writer.close()

        # Apply formatting using openpyxl
        self.apply_styles(output_path)

        # Generate sample radar charts for top companies
        for grp in top_groups[:2]:
            grp_companies = df_peers[df_peers["effective_peer_group"] == grp]
            if len(grp_companies) > 0:
                top_company = grp_companies.sort_values(by="overall_peer_percentile", ascending=False).iloc[0]
                categories = ["ROE", "ROCE", "Low Debt", "NPM", "OPM", "FCF"]
                vals = [
                    float(top_company.get("return_on_equity_pct_percentile", 50)),
                    float(top_company.get("roce_pct_percentile", 50)),
                    float(top_company.get("debt_to_equity_percentile", 50)),
                    float(top_company.get("net_profit_margin_pct_percentile", 50)),
                    float(top_company.get("operating_profit_margin_pct_percentile", 50)),
                    float(top_company.get("free_cash_flow_cr_percentile", 50))
                ]
                generate_peer_radar_chart(top_company["company_name"], categories, vals)

        logger.info(f"Successfully generated Peer Comparison report at {output_path}")
        return output_path

    def apply_styles(self, file_path: str) -> None:
        """Applies header formatting, gridlines, auto-column widths, and percentile shading."""
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.load_workbook(file_path)

        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for sheet in wb.worksheets:
            sheet.views.sheetView[0].showGridLines = True
            
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.border = thin_border
                    cell.font = Font(name="Calibri", size=10)

                    col_name = str(sheet.cell(row=1, column=cell.column).value).lower()
                    val = cell.value

                    if isinstance(val, (int, float)) and "percentile" in col_name and val >= 75.0:
                        cell.fill = green_fill

            for col in sheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(file_path)
