"""
Screener Excel Exporter Module.
Exports screener results across all 6 presets into a multi-tab formatted Excel report.
"""

import os
import logging
from typing import Optional
import pandas as pd

from src.analytics.screener_engine import ScreenerEngine
from src.analytics.presets import ALL_PRESETS

logger = logging.getLogger(__name__)


class ScreenerExporter:
    """Exports screener evaluation output to Excel workbooks."""

    def __init__(
        self, engine: Optional[ScreenerEngine] = None, output_dir: str = "output"
    ):
        self.engine = engine or ScreenerEngine()
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def export_all_presets_to_excel(
        self, file_name: str = "screener_results.xlsx"
    ) -> str:
        """
        Executes all 6 presets, calculates composite scores and sector ranks,
        and writes a multi-sheet Excel file.
        """
        output_path = os.path.join(self.output_dir, file_name)
        self.engine.load_latest_company_ratios()

        summary_rows = []
        excel_writer = pd.ExcelWriter(output_path, engine="openpyxl")

        for preset_key, config in ALL_PRESETS.items():
            self.engine.set_config(config)
            df_filtered = self.engine.apply_filters()
            df_scored = self.engine.compute_sector_relative_scores(df_filtered)

            # Sort by composite score descending
            if "composite_score" in df_scored.columns:
                df_scored = df_scored.sort_values(by="composite_score", ascending=False)

            sheet_name = config.name[:31]  # Excel sheet name length limit

            # Select key columns for export
            export_cols = [
                "company_id",
                "company_name",
                "broad_sector",
                "sub_sector",
                "composite_score",
                "sector_relative_score",
                "sector_rank",
                "return_on_equity_pct",
                "roce_pct",
                "debt_to_equity",
                "net_profit_margin_pct",
                "operating_profit_margin_pct",
                "free_cash_flow_cr",
                "revenue_cagr_5yr",
                "pat_cagr_5yr",
                "dividend_payout_ratio_pct",
            ]
            existing_cols = [c for c in export_cols if c in df_scored.columns]
            export_df = df_scored[existing_cols]

            export_df.to_excel(excel_writer, sheet_name=sheet_name, index=False)

            summary_rows.append(
                {
                    "Preset Key": preset_key,
                    "Preset Name": config.name,
                    "Description": config.description,
                    "Matching Companies Count": len(df_scored),
                    "Top Company": (
                        df_scored.iloc[0]["company_name"]
                        if len(df_scored) > 0
                        else "N/A"
                    ),
                    "Top Composite Score": (
                        df_scored.iloc[0]["composite_score"]
                        if len(df_scored) > 0
                        else 0.0
                    ),
                }
            )

        df_summary = pd.DataFrame(summary_rows)
        df_summary.to_excel(excel_writer, sheet_name="Summary Overview", index=False)

        excel_writer.close()

        # Apply formatting styles using openpyxl
        self.apply_openpyxl_styles(output_path)

        logger.info(
            f"Successfully exported and formatted Screener results to {output_path}"
        )
        return output_path

    def apply_openpyxl_styles(self, file_path: str) -> None:
        """Applies headers, auto column width, and conditional color highlights to Excel sheets."""
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.load_workbook(file_path)

        header_fill = PatternFill(
            start_color="1F497D", end_color="1F497D", fill_type="solid"
        )
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        green_fill = PatternFill(
            start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
        )
        red_fill = PatternFill(
            start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
        )

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for sheet in wb.worksheets:
            sheet.views.sheetView[0].showGridLines = True

            # Format Headers
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Data Rows Styling
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.border = thin_border
                    cell.font = Font(name="Calibri", size=10)

                    # Conditional Highlights based on header column
                    col_name = str(sheet.cell(row=1, column=cell.column).value).lower()
                    val = cell.value

                    if isinstance(val, (int, float)):
                        if (
                            "roe" in col_name
                            or "roce" in col_name
                            or "score" in col_name
                        ):
                            if val >= 15.0 or val >= 75.0:
                                cell.fill = green_fill
                        elif "debt_to_equity" in col_name:
                            if val > 1.0:
                                cell.fill = red_fill
                            elif val <= 0.5:
                                cell.fill = green_fill

            # Auto Column Widths
            for col in sheet.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(file_path)
